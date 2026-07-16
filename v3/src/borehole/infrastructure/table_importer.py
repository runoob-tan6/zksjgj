"""Excel/CSV 钻孔统计表导入器。

基于 generate-borehole-files skill 的逻辑：
- 行1: 钻孔结构（每列一个值）
- 行2: 风化代码
- 行3: 图案/代码（用于 .-c）
- 行4: 地层时代（用于 .-b）
- 行5: 岩土名称（用于 .-h）
- 行6起: 数据行，B列=编号，C列=高程，D列起=各层底深度
"""

from __future__ import annotations

import re
from datetime import date, timedelta
from math import isfinite
from pathlib import Path

from ..domain.models import ProjectData
from ..application.project_service import load_project

WATER = "水位"
TIME = "时间"
DEFAULT_LOCATION = "八曲河"
DEFAULT_STAGE = "初步设计"
DEFAULT_DESIGNER = "长沙核工业工程勘察院有限公司"
DEFAULT_PROJECT = "宁乡市八曲河河道治理工程"
DEFAULT_START_DATE = "2026.5.10"
DEFAULT_HOLES_PER_DAY = 2


def _fmt(value) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, float):
        return f"{value:g}"
    return str(value).strip()


def _parse_date(value: str) -> date:
    parts = re.split(r"[.\-/]", value.strip())
    if len(parts) != 3:
        raise ValueError(f"日期格式错误：{value!r}，应为 YYYY.M.D")
    return date(int(parts[0]), int(parts[1]), int(parts[2]))


def _make_text(lines: list[str]) -> str:
    return "\n".join([*lines, "★"])


def _pair_text(pairs: list[tuple[str, str]]) -> str:
    return _make_text([f"{depth},{value}" for depth, value in pairs if depth and value])


def _h_text(layers: list[dict]) -> str:
    lines: list[str] = []
    for layer in layers:
        if layer["depth"] or layer["name"]:
            lines.append(f"#{layer['depth']}")
            lines.append(layer["name"])
    return _make_text(lines)


def _compressed_pairs(layers: list[dict], field: str) -> list[tuple[str, str]]:
    deepest: dict[str, tuple[float, str]] = {}
    for layer in layers:
        value = layer.get(field, "")
        if not value:
            continue
        current = deepest.get(value)
        if current is None or layer["depth_num"] > current[0]:
            deepest[value] = (layer["depth_num"], layer["depth"])
    return [
        (depth, value)
        for value, (_depth_num, depth) in sorted(deepest.items(), key=lambda item: item[1][0])
    ]


def _render_main_file(
    hole_id: str,
    depth: str,
    elevation: str,
    work_date: str,
    project: str,
    location: str,
    stage: str,
    designer: str,
) -> str:
    return _make_text([
        hole_id,
        depth,
        elevation,
        location,
        ",90",
        "100",
        work_date,
        project,
        "001",
        stage,
        "0,0",
        work_date,
        depth,
        "L",
        ",90",
        designer,
    ])


def _try_openpyxl(path: Path):
    try:
        from openpyxl import load_workbook
        return load_workbook(path, data_only=True)
    except ImportError:
        return None


def _discover_layer_headers(ws) -> list[dict]:
    headers: list[dict] = []
    for col in range(4, ws.max_column + 1):
        structure = _fmt(ws.cell(1, col).value)
        weathering = _fmt(ws.cell(2, col).value)
        code = _fmt(ws.cell(3, col).value)
        formation = _fmt(ws.cell(4, col).value)
        name = _fmt(ws.cell(5, col).value)
        if name in {WATER, TIME} or formation in {WATER, TIME} or code in {WATER, TIME}:
            continue
        if name or code:
            headers.append({
                "col": col,
                "structure": structure,
                "weathering": weathering,
                "code": code,
                "formation": formation,
                "name": name,
            })
    return headers


def _read_layers(ws, row: int, headers: list[dict]) -> list[dict]:
    layers: list[dict] = []
    for header in headers:
        depth = _fmt(ws.cell(row, header["col"]).value)
        if not depth:
            continue
        try:
            depth_num = float(depth)
        except ValueError:
            continue
        if not isfinite(depth_num):
            raise ValueError(f"层底深度必须是有限数字：{depth}")
        layers.append({
            "depth": depth,
            "depth_num": depth_num,
            "code": header["code"],
            "formation": header["formation"],
            "structure": header["structure"],
            "weathering": header["weathering"],
            "name": header["name"],
        })
    layers.sort(key=lambda l: l["depth_num"])
    return layers


def _output_hole_id(source_hole: str, output_prefix: str) -> str:
    match = re.search(r"(\d+)$", source_hole)
    if not match:
        return source_hole if not output_prefix else f"{output_prefix}{source_hole}"
    if not output_prefix:
        return source_hole
    return f"{output_prefix}{match.group(1)}"


def import_from_table(
    file_path: Path,
    project_folder: Path,
    output_prefix: str = "",
    start_date: str = DEFAULT_START_DATE,
    holes_per_day: int = DEFAULT_HOLES_PER_DAY,
    project_name: str = "",
    location: str = DEFAULT_LOCATION,
    stage: str = DEFAULT_STAGE,
    designer: str = DEFAULT_DESIGNER,
) -> ProjectData:
    """从 Excel/CSV 统计表导入钻孔数据到项目文件夹。"""
    if file_path.suffix.lower() in (".xlsx", ".xls"):
        wb = _try_openpyxl(file_path)
        if not wb:
            raise ImportError("需要 openpyxl 库来读取 Excel 文件。")
        ws = wb.active
        headers = _discover_layer_headers(ws)
        if not headers:
            raise ValueError("未找到岩性层数据列。")
        proj_name = project_name or _fmt(ws.cell(1, 1).value) or DEFAULT_PROJECT
        first_date = _parse_date(start_date)

        project_folder.mkdir(parents=True, exist_ok=True)
        imported = 0

        for row in range(6, ws.max_row + 1):
            source_hole = _fmt(ws.cell(row, 2).value)
            if not source_hole:
                continue
            match = re.search(r"(\d+)$", source_hole)
            if not match:
                continue
            hole_number = int(match.group(1))
            hole_id = _output_hole_id(source_hole, output_prefix)
            elevation = _fmt(ws.cell(row, 3).value)
            layers = _read_layers(ws, row, headers)
            if not layers:
                continue

            depth = layers[-1]["depth"]
            work_date = first_date + timedelta(days=(hole_number - 1) // holes_per_day)
            work_date_text = f"{work_date.year}.{work_date.month}.{work_date.day}"

            files = {
                hole_id: _render_main_file(
                    hole_id=hole_id,
                    depth=depth,
                    elevation=elevation,
                    work_date=work_date_text,
                    project=proj_name,
                    location=location,
                    stage=stage,
                    designer=designer,
                ),
                f"{hole_id}.-b": _pair_text(_compressed_pairs(layers, "formation")),
                f"{hole_id}.-c": _pair_text([(l["depth"], l["code"]) for l in layers]),
                f"{hole_id}.-d": _pair_text(_compressed_pairs(layers, "structure")),
                f"{hole_id}.-g": _pair_text(_compressed_pairs(layers, "weathering")),
                f"{hole_id}.-h": _h_text(layers),
            }

            for name, text in files.items():
                target = project_folder / name
                if target.exists():
                    _backup_with_encoding(target, project_folder)
                target.write_text(text, encoding="utf-8")
            imported += 1

    elif file_path.suffix.lower() == ".csv":
        import csv
        with file_path.open(encoding="utf-8-sig") as f:
            reader = csv.reader(f)
            rows = [list(row) for row in reader]
        return _import_csv_rows(rows, project_folder, output_prefix, start_date, holes_per_day,
                                project_name, location, stage, designer)
    else:
        raise ValueError(f"不支持的文件格式：{file_path.suffix}")

    return load_project(project_folder)


def _import_csv_rows(
    rows: list[list[str]],
    project_folder: Path,
    output_prefix: str,
    start_date: str,
    holes_per_day: int,
    project_name: str,
    location: str,
    stage: str,
    designer: str,
) -> ProjectData:
    if len(rows) < 6:
        raise ValueError("CSV 行数不足，无法解析。")

    headers: list[dict] = []
    for j in range(3, len(rows[0])):
        structure = rows[0][j].strip() if j < len(rows[0]) else ""
        weathering = rows[1][j].strip() if j < len(rows[1]) else ""
        code = rows[2][j].strip() if j < len(rows[2]) else ""
        formation = rows[3][j].strip() if j < len(rows[3]) else ""
        name = rows[4][j].strip() if j < len(rows[4]) else ""
        if name in {WATER, TIME} or formation in {WATER, TIME} or code in {WATER, TIME}:
            continue
        if name or code:
            headers.append({
                "col": j,
                "structure": structure,
                "weathering": weathering,
                "code": code,
                "formation": formation,
                "name": name,
            })

    if not headers:
        raise ValueError("未找到岩性层数据列。")

    proj_name = project_name or (rows[0][0].strip() if rows[0] else "") or DEFAULT_PROJECT
    first_date = _parse_date(start_date)
    project_folder.mkdir(parents=True, exist_ok=True)
    imported = 0

    for i in range(5, len(rows)):
        row = rows[i]
        if len(row) < 3:
            continue
        source_hole = row[1].strip()
        if not source_hole:
            continue
        match = re.search(r"(\d+)$", source_hole)
        if not match:
            continue
        hole_number = int(match.group(1))
        hole_id = _output_hole_id(source_hole, output_prefix)
        elevation = row[2].strip()

        layers: list[dict] = []
        for header in headers:
            col = header["col"]
            if col >= len(row):
                continue
            depth = row[col].strip()
            if not depth:
                continue
            try:
                depth_num = float(depth)
            except ValueError:
                continue
            if not isfinite(depth_num):
                raise ValueError(f"层底深度必须是有限数字：{depth}")
            layers.append({
                "depth": depth,
                "depth_num": depth_num,
                "code": header["code"],
                "formation": header["formation"],
                "structure": header["structure"],
                "weathering": header["weathering"],
                "name": header["name"],
            })
        layers.sort(key=lambda l: l["depth_num"])
        if not layers:
            continue

        depth = layers[-1]["depth"]
        work_date = first_date + timedelta(days=(hole_number - 1) // holes_per_day)
        work_date_text = f"{work_date.year}.{work_date.month}.{work_date.day}"

        files = {
            hole_id: _render_main_file(
                hole_id=hole_id,
                depth=depth,
                elevation=elevation,
                work_date=work_date_text,
                project=proj_name,
                location=location,
                stage=stage,
                designer=designer,
            ),
            f"{hole_id}.-b": _pair_text(_compressed_pairs(layers, "formation")),
            f"{hole_id}.-c": _pair_text([(l["depth"], l["code"]) for l in layers]),
            f"{hole_id}.-d": _pair_text(_compressed_pairs(layers, "structure")),
            f"{hole_id}.-g": _pair_text(_compressed_pairs(layers, "weathering")),
            f"{hole_id}.-h": _h_text(layers),
        }

        for name, text in files.items():
            target = project_folder / name
            if target.exists():
                _backup_with_encoding(target, project_folder)
            target.write_text(text, encoding="utf-8")
        imported += 1

    return load_project(project_folder)


def _detect_file_encoding(path: Path) -> str:
    """检测文件编码。"""
    for encoding in ("utf-8", "gbk", "ansi"):
        try:
            path.read_text(encoding=encoding)
            return encoding
        except (UnicodeDecodeError, UnicodeError):
            continue
    return "utf-8"


def _backup_with_encoding(file_path: Path, project_folder: Path) -> None:
    """备份文件，使用编码检测和带时间戳的文件名。"""
    from datetime import datetime
    backup_dir = project_folder / "tmp"
    backup_dir.mkdir(exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = backup_dir / f"{file_path.name}.{timestamp}.bak"
    encoding = _detect_file_encoding(file_path)
    content = file_path.read_text(encoding=encoding)
    backup.write_text(content, encoding=encoding)
