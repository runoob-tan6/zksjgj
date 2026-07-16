"""XLSX 导出 - 使用 openpyxl，统计汇总和标贯分析使用 Excel 公式。"""

from __future__ import annotations

import re
from collections.abc import Iterator
from math import isfinite
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from ..domain.models import BasicLayer, Borehole, ProjectData, TestRecord

LAYER_TEST_TYPES: dict[str, str] = {
    "o": "取样",
    "q": "标贯",
    "n": "注水",
    "m": "压水",
}

LAYER_TEST_VALUE_NAMES: dict[str, str] = {
    "o": "样品编号",
    "q": "标贯击数",
    "n": "渗透系数",
    "m": "透水率",
}

EXPORT_HEADERS = [
    "钻孔编号",
    "层序号",
    "地层时代/成因",
    "岩性代号",
    "试验类型",
    "试验深度",
    "结果值/编号",
    "结果字段",
]

SUMMARY_HEADERS = [
    "地层时代/成因",
    "岩性代号",
    "试验类型",
    "试验数",
    "最小值",
    "最大值",
    "平均值",
]


def _to_float(value: str) -> float | None:
    try:
        number = float(str(value).strip())
    except (TypeError, ValueError):
        return None
    return number if isfinite(number) else None


def _fmt_depth(value: float, suffix: str = "") -> str:
    if suffix == "q":
        return f"{value:.2f}"
    return f"{value:.1f}"


def _fmt_test_result(suffix: str, value: str) -> str:
    if suffix == "n":
        number = _to_float(value)
        if number is None:
            return value
        return f"{number:.2E}"
    if suffix in ("o", "q"):
        return value
    number = _to_float(value)
    if number is None:
        return value
    return f"{number:.1f}"


def _borehole_sort_key(prefix: str) -> tuple[int, str | int]:
    match = re.search(r"(\d+)$", prefix)
    if match:
        return (0, int(match.group(1)))
    return (1, prefix)


def _layer_ranges(borehole: Borehole) -> Iterator[tuple[int, float, float, BasicLayer]]:
    top = 0.0
    for index, layer in enumerate(borehole.layers, start=1):
        bottom = _to_float(layer.bottom_depth)
        if bottom is None or bottom <= top:
            continue
        yield index, top, bottom, layer
        top = bottom


def _test_matches_layer(
    suffix: str, test_top: float, test_bottom: float, layer_top: float, layer_bottom: float
) -> bool:
    return layer_top <= test_top < layer_bottom


def _effective_layer_formation(borehole: Borehole, layer_index: int, layer: BasicLayer) -> str:
    if layer.formation:
        return layer.formation
    for next_layer in borehole.layers[layer_index:]:
        if next_layer.formation:
            return next_layer.formation
    return ""


def _layer_test_rows(boreholes: list[Borehole]) -> list[list[Any]]:
    rows: list[list[Any]] = []
    for borehole in boreholes:
        ranges = list(_layer_ranges(borehole))
        if not ranges:
            continue
        for suffix, test_type in LAYER_TEST_TYPES.items():
            for record in borehole.tests.get(suffix, []):
                values = [str(value).strip() for value in record.values]
                while len(values) < 3:
                    values.append("")
                test_top = _to_float(values[0])
                test_bottom = _to_float(values[1])
                if test_top is None or test_bottom is None or test_bottom <= test_top:
                    continue
                result_value = values[2]
                for layer_index, layer_top, layer_bottom, layer in ranges:
                    if not _test_matches_layer(suffix, test_top, test_bottom, layer_top, layer_bottom):
                        continue
                    formation = _effective_layer_formation(borehole, layer_index, layer)
                    rows.append([
                        borehole.prefix,
                        layer_index,
                        formation,
                        layer.lithology_code,
                        test_type,
                        f"{_fmt_depth(test_top, suffix)}-{_fmt_depth(test_bottom, suffix)}",
                        _fmt_test_result(suffix, result_value),
                        LAYER_TEST_VALUE_NAMES[suffix],
                    ])
                    if suffix != "m":
                        break
    return rows


def _get_unique_groups(rows: list[list]) -> list[tuple[str, str, str]]:
    seen: set[tuple[str, str, str]] = set()
    groups: list[tuple[str, str, str]] = []
    for row in rows:
        key = (str(row[2]), str(row[3]), str(row[4]))
        if key not in seen:
            seen.add(key)
            groups.append(key)
    return groups


def _alpha_formula(rod_length_cell: str) -> str:
    """生成杆长修正系数的 Excel IF 公式。"""
    c = rod_length_cell
    return (
        f'IF({c}<=3,1,'
        f'IF({c}<=6,0.92,'
        f'IF({c}<=9,0.86,'
        f'IF({c}<=12,0.81,'
        f'IF({c}<=15,0.77,'
        f'IF({c}<=18,0.73,'
        f'IF({c}<=21,0.7,"超长")))))))'
    )


def _auto_width(ws: Any, max_width: int = 40) -> None:
    """自动调整列宽，跳过公式单元格。"""
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            val = cell.value
            if val and not (isinstance(val, str) and val.startswith("=")):
                max_length = max(max_length, len(str(val)))
        ws.column_dimensions[column_letter].width = min(max(max_length + 3, 8), max_width)


def _write_xlsx(path: Path, rows: list[list[Any]], boreholes: list[Borehole]) -> None:
    wb = Workbook()
    header_font = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center")

    # ── Sheet 1: 试验汇总 ──
    ws = wb.active
    if ws is None:
        raise RuntimeError("无法创建试验汇总工作表。")
    ws.title = "试验汇总"

    for col_idx, header in enumerate(EXPORT_HEADERS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.alignment = center

    data_end_row = len(rows) + 1
    for row_idx, row_data in enumerate(rows, 2):
        for col_idx, value in enumerate(row_data, 1):
            data_cell = ws.cell(row=row_idx, column=col_idx)
            if col_idx == 7:
                result_val = value
                if isinstance(result_val, str):
                    number = _to_float(result_val)
                    data_cell.value = number if number is not None else result_val
                else:
                    data_cell.value = result_val
                if row_data[4] == "注水" and isinstance(data_cell.value, float):
                    data_cell.number_format = "0.00E+00"
            elif isinstance(value, (int, float)):
                data_cell.value = value
            else:
                data_cell.value = value

    # ── 统计汇总（使用 Excel 公式）──
    groups = _get_unique_groups(rows)
    summary_start = data_end_row + 2
    ws.cell(row=summary_start, column=1, value="统计汇总").font = header_font
    for col_idx, header in enumerate(SUMMARY_HEADERS, 1):
        cell = ws.cell(row=summary_start + 1, column=col_idx, value=header)
        cell.font = header_font
        cell.alignment = center

    # 数据列引用：C=地层时代, D=岩性代号, E=试验类型, G=结果值
    for i, (formation, lithology, test_type) in enumerate(groups):
        r = summary_start + 2 + i
        ws.cell(row=r, column=1, value=formation)
        ws.cell(row=r, column=2, value=lithology)
        ws.cell(row=r, column=3, value=test_type)

        f_formation = f"$C$2:$C${data_end_row}"
        f_lithology = f"$D$2:$D${data_end_row}"
        f_type = f"$E$2:$E${data_end_row}"
        f_result = f"$G$2:$G${data_end_row}"
        criteria_f = f'"{formation}"'
        criteria_l = f'"{lithology}"'
        criteria_t = f'"{test_type}"'

        # 试验数
        ws.cell(row=r, column=4).value = (
            f'=COUNTIFS({f_formation},{criteria_f},{f_lithology},{criteria_l},{f_type},{criteria_t})'
        )
        # 最小值
        ws.cell(row=r, column=5).value = (
            f'=IFERROR(MINIFS({f_result},{f_formation},{criteria_f},{f_lithology},{criteria_l},{f_type},{criteria_t}),"")'
        )
        # 最大值
        ws.cell(row=r, column=6).value = (
            f'=IFERROR(MAXIFS({f_result},{f_formation},{criteria_f},{f_lithology},{criteria_l},{f_type},{criteria_t}),"")'
        )
        # 平均值
        ws.cell(row=r, column=7).value = (
            f'=IFERROR(AVERAGEIFS({f_result},{f_formation},{criteria_f},{f_lithology},{criteria_l},{f_type},{criteria_t}),"")'
        )

        if test_type == "注水":
            for c in (5, 6, 7):
                ws.cell(row=r, column=c).number_format = "0.00E+00"

    _auto_width(ws)

    # ── Sheet 2: 钻孔汇总 ──
    ws2 = wb.create_sheet("钻孔汇总")
    borehole_headers = [
        "钻孔编号",
        "孔口高程(m)",
        "深度(m)",
        "勘探开始日期",
        "勘探结束日期",
        "地下水埋深(m)",
        "水位观测日期",
    ]
    for col_idx, header in enumerate(borehole_headers, 1):
        cell = ws2.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.alignment = center

    total_depth = 0.0
    for row_idx, borehole in enumerate(boreholes, 2):
        lines = borehole.main.normalized_lines()
        hole_id = lines[0]
        elevation_str = lines[2]
        depth_str = lines[1]
        start_date = lines[6]
        end_date = lines[11]

        depth_number = _to_float(depth_str)
        if depth_number is not None:
            depth_val: float | str = depth_number
            total_depth += depth_number
        else:
            depth_val = depth_str

        ws2.cell(row=row_idx, column=1, value=hole_id)
        cell_elev = ws2.cell(row=row_idx, column=2)
        elevation_number = _to_float(elevation_str)
        if elevation_number is not None:
            cell_elev.value = elevation_number
            cell_elev.number_format = "0.00"
        else:
            cell_elev.value = elevation_str
        cell_depth = ws2.cell(row=row_idx, column=3)
        if isinstance(depth_val, float):
            cell_depth.value = depth_val
            cell_depth.number_format = "0.0"
        else:
            cell_depth.value = depth_val
        ws2.cell(row=row_idx, column=4, value=start_date)
        ws2.cell(row=row_idx, column=5, value=end_date)

        water_records = borehole.tests.get("l", [])
        if water_records:
            first = water_records[0]
            if len(first.values) > 0 and first.values[0].strip():
                cell_wl = ws2.cell(row=row_idx, column=6)
                water_number = _to_float(first.values[0])
                if water_number is not None:
                    cell_wl.value = water_number
                    cell_wl.number_format = "0.0"
                else:
                    cell_wl.value = first.values[0]
            if len(first.values) > 1 and first.values[1].strip():
                ws2.cell(row=row_idx, column=7, value=first.values[1])

    total_row = len(boreholes) + 2
    ws2.cell(row=total_row, column=1, value="总深度").font = header_font
    cell_total = ws2.cell(row=total_row, column=3)
    cell_total.value = total_depth
    cell_total.number_format = "0.0"
    cell_total.font = header_font

    _auto_width(ws2, 30)

    # ── Sheet 3: 标贯分析 ──
    ws3 = wb.create_sheet("标贯分析")

    # 收集全部标贯记录
    spt_records: list[tuple[str, Borehole, TestRecord]] = []
    for borehole in boreholes:
        for record in borehole.tests.get("q", []):
            spt_records.append((borehole.prefix, borehole, record))

    if not spt_records:
        wb.save(path)
        return

    # 修正明细表
    detail_headers = ["钻孔编号", "起始深度", "终止深度", "杆长(m)", "实测N", "修正系数α", "修正N'"]
    for col_idx, header in enumerate(detail_headers, 1):
        cell = ws3.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.alignment = center

    detail_end_row = len(spt_records) + 1
    for row_idx, (prefix, _borehole, record) in enumerate(spt_records, 2):
        values = list(record.values)
        while len(values) < 3:
            values.append("")
        start_depth = values[0].strip()
        end_depth = values[1].strip()
        raw_n_str = values[2].strip()

        rod_length = _to_float(start_depth) or 0.0
        raw_n = _to_float(raw_n_str) or 0.0

        ws3.cell(row=row_idx, column=1, value=prefix)
        ws3.cell(row=row_idx, column=2, value=start_depth)
        ws3.cell(row=row_idx, column=3, value=end_depth)
        cell_rod = ws3.cell(row=row_idx, column=4)
        cell_rod.value = rod_length
        cell_rod.number_format = "0.00"
        cell_raw = ws3.cell(row=row_idx, column=5)
        cell_raw.value = raw_n

        # 修正系数 α 公式
        alpha_cell = ws3.cell(row=row_idx, column=6)
        alpha_cell.value = f"={_alpha_formula(f'D{row_idx}')}"
        alpha_cell.number_format = "0.00"

        # 修正 N' = α × N 公式
        corrected_cell = ws3.cell(row=row_idx, column=7)
        corrected_cell.value = f'=IF(F{row_idx}="超长","超长",F{row_idx}*E{row_idx})'
        corrected_cell.number_format = "0.00"

    # 收集地层列表
    layer_codes: list[str] = []
    seen_codes: set[str] = set()
    for _prefix, borehole, _record in spt_records:
        for layer in borehole.layers:
            code = layer.lithology_code.strip()
            if code and code not in seen_codes:
                seen_codes.add(code)
                layer_codes.append(code)

    # 分层统计（使用 Excel 公式）
    stats_start = detail_end_row + 2
    ws3.cell(row=stats_start, column=1, value="分层统计参数").font = header_font

    stats_headers = [
        "地层", "样本数n", "最小值", "最大值", "平均值Φm",
        "标准差σf", "变异系数δ", "统计修正系数γs", "标准值Φk", "备注"
    ]
    for col_idx, header in enumerate(stats_headers, 1):
        cell = ws3.cell(row=stats_start + 1, column=col_idx, value=header)
        cell.font = header_font
        cell.alignment = center

    # A列=钻孔编号, G列=修正N'
    # 用 COUNTIF/AVERAGEIF 等按钻孔编号前缀无法直接按地层分组
    # 改用辅助列：H列=地层代号（从钻孔的 layers 匹配）
    # 但 openpyxl 无法写复杂的跨表匹配公式，改用直接按地层名称统计
    # 用 COUNTIF 匹配 H 列的地层代号

    # 先写辅助列头
    ws3.cell(row=1, column=8, value="地层代号").font = header_font

    # 为每条记录填入地层代号（根据深度匹配）
    for row_idx, (prefix, borehole, record) in enumerate(spt_records, 2):
        values = record.values
        start_depth_str = values[0].strip() if values else ""
        depth = _to_float(start_depth_str) or 0.0
        layer_code = ""
        for layer in borehole.layers:
            ld = _to_float(layer.bottom_depth)
            if ld is None:
                continue
            if depth < ld:
                layer_code = layer.lithology_code
                break
        if not layer_code and borehole.layers:
            layer_code = borehole.layers[-1].lithology_code
        ws3.cell(row=row_idx, column=8, value=layer_code)

    # 分层统计公式行
    for i, code in enumerate(layer_codes):
        r = stats_start + 2 + i
        ws3.cell(row=r, column=1, value=code)

        h_range = f"$H$2:$H${detail_end_row}"
        g_range = f"$G$2:$G${detail_end_row}"
        criteria = f'"{code}"'

        # n
        ws3.cell(row=r, column=2).value = f'=COUNTIF({h_range},{criteria})'
        # 最小值
        ws3.cell(row=r, column=3).value = f'=IFERROR(MINIFS({g_range},{h_range},{criteria}),"")'
        # 最大值
        ws3.cell(row=r, column=4).value = f'=IFERROR(MAXIFS({g_range},{h_range},{criteria}),"")'
        # 平均值 Φm
        ws3.cell(row=r, column=5).value = f'=IFERROR(ROUND(AVERAGEIF({h_range},{criteria},{g_range}),3),"")'
        # 标准差 σf (样本标准差，用 SUMPRODUCT 替代数组公式)
        ws3.cell(row=r, column=6).value = (
            f'=IFERROR(IF(B{r}>=2,ROUND(SQRT(SUMPRODUCT(({h_range}={criteria})*({g_range}-E{r})^2)/(B{r}-1)),3),""),"")'
        )
        # 变异系数 δ = σf / Φm
        ws3.cell(row=r, column=7).value = f'=IFERROR(IF(F{r}<>"",ROUND(F{r}/E{r},3),""),"")'
        # 统计修正系数 γs = 1 - (1.704/√n + 4.678/n²) × δ
        ws3.cell(row=r, column=8).value = (
            f'=IFERROR(IF(B{r}>=3,ROUND(1-(1.704/SQRT(B{r})+4.678/B{r}^2)*G{r},3),""),"")'
        )
        # 标准值 Φk = γs × Φm
        ws3.cell(row=r, column=9).value = f'=IFERROR(IF(H{r}<>"",ROUND(H{r}*E{r},3),""),"")'
        # 备注
        ws3.cell(row=r, column=10).value = (
            f'=IF(B{r}<3,"样本不足",IF(B{r}<6,"仅供参考",IF(G{r}>0.3,"变异性较高","")))'
        )

    # 列宽 + 居中对齐
    center = Alignment(horizontal="center", vertical="center")
    for ws_item in [ws, ws2, ws3]:
        for col in ws_item.column_dimensions:
            ws_item.column_dimensions[col].width = 15
        for row in ws_item.iter_rows():
            for cell in row:
                cell.alignment = center

    wb.save(path)


def _write_csv(path: Path, rows: list[list[Any]], boreholes: list[Borehole]) -> None:
    import csv

    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(EXPORT_HEADERS)
        writer.writerows(rows)
        groups = _get_unique_groups(rows)
        if groups:
            writer.writerow([])
            writer.writerow(["统计汇总"])
            writer.writerow(SUMMARY_HEADERS)
            for formation, lithology, test_type in groups:
                count = 0
                values = []
                for row in rows:
                    if str(row[2]) == formation and str(row[3]) == lithology and str(row[4]) == test_type:
                        count += 1
                        number = _to_float(str(row[6]))
                        if number is not None:
                            values.append(number)
                if values:
                    writer.writerow([formation, lithology, test_type, count,
                                     min(values), max(values), sum(values) / len(values)])
                else:
                    writer.writerow([formation, lithology, test_type, count, "", "", ""])
        writer.writerow([])
        writer.writerow(["钻孔汇总"])
        writer.writerow(
            ["钻孔编号", "孔口高程(m)", "深度(m)", "勘探开始日期", "勘探结束日期", "地下水埋深(m)", "水位观测日期"]
        )
        total_depth = 0.0
        for borehole in boreholes:
            lines = borehole.main.normalized_lines()
            hole_id = lines[0]
            elevation_str = lines[2]
            depth_str = lines[1]
            start_date = lines[6]
            end_date = lines[11]
            depth_number = _to_float(depth_str)
            if depth_number is not None:
                depth_val: float | str = depth_number
                total_depth += depth_number
            else:
                depth_val = depth_str
            water_depth = ""
            water_date = ""
            water_records = borehole.tests.get("l", [])
            if water_records:
                first = water_records[0]
                if len(first.values) > 0 and first.values[0].strip():
                    water_depth = first.values[0]
                if len(first.values) > 1 and first.values[1].strip():
                    water_date = first.values[1]
            writer.writerow([hole_id, elevation_str, depth_val, start_date, end_date, water_depth, water_date])
        writer.writerow(["总深度", "", total_depth, "", "", "", ""])


def export_layer_test_summary(project: ProjectData, target_path: Path) -> int:
    """导出地层试验汇总表，支持 .xlsx 和 .csv。"""
    boreholes = project.sorted_boreholes()
    rows = _layer_test_rows(boreholes)
    target_path.parent.mkdir(parents=True, exist_ok=True)

    if target_path.suffix.lower() == ".xlsx":
        _write_xlsx(target_path, rows, boreholes)
    else:
        _write_csv(target_path, rows, boreholes)
    return len(rows)
