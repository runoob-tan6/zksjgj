from pathlib import Path

import pytest
from openpyxl import Workbook

from borehole.infrastructure.table_importer import import_from_table


@pytest.mark.parametrize("value", ["nan", "NaN", "inf", "-inf"])
def test_import_rejects_non_finite_depth(tmp_path: Path, value: str) -> None:
    source = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.cell(1, 1, "测试项目")
    sheet.cell(1, 4, "0.11")
    sheet.cell(2, 4, "f")
    sheet.cell(3, 4, "11")
    sheet.cell(4, 4, "Q4")
    sheet.cell(5, 4, "粉质黏土")
    sheet.cell(6, 2, "ZK1")
    sheet.cell(6, 3, 100.0)
    sheet.cell(6, 4, value)
    workbook.save(source)

    with pytest.raises(ValueError, match="有限数字"):
        import_from_table(source, tmp_path / "output")
