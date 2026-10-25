from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import Workbook

from borehole.infrastructure.table_importer import import_from_table


def test_import_uses_xlsx_time_column_for_start_and_end_dates(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.cell(1, 1, "测试项目")
    sheet.cell(1, 4, "0.11")
    sheet.cell(2, 4, "f")
    sheet.cell(3, 4, "11")
    sheet.cell(4, 4, "Q4")
    sheet.cell(5, 4, "粉质黏土")
    sheet.cell(5, 5, "时间")
    sheet.cell(6, 2, "ZK1")
    sheet.cell(6, 3, 100.0)
    sheet.cell(6, 4, 5.0)
    sheet.cell(6, 5, datetime(2023, 11, 5))
    workbook.save(source)

    import_from_table(source, tmp_path / "output")

    main_lines = (tmp_path / "output" / "ZK1").read_text(encoding="utf-8").splitlines()
    assert main_lines[6] == "2023.11.5"
    assert main_lines[11] == "2023.11.5"


def test_import_uses_csv_time_column_for_start_and_end_dates(tmp_path: Path) -> None:
    source = tmp_path / "source.csv"
    source.write_text(
        "测试项目,,,\n"
        ",,,\n"
        ",,,11,\n"
        ",,,Q4,\n"
        ",,,粉质黏土,时间\n"
        ",ZK1,100,5,2023.11.6\n",
        encoding="utf-8",
    )

    import_from_table(source, tmp_path / "output")

    main_lines = (tmp_path / "output" / "ZK1").read_text(encoding="utf-8").splitlines()
    assert main_lines[6] == "2023.11.6"
    assert main_lines[11] == "2023.11.6"


@pytest.mark.parametrize("value", ["nan", "NaN", "inf", "-inf"])
def test_import_rejects_non_finite_depth(tmp_path: Path, value: str) -> None:
    source = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.cell(1, 1, "测试项目")
    sheet.cell(1, 4, "0.11")
    sheet.cell(2, 4, "f")
    sheet.cell(3, 4, "11")
    sheet.cell(4, 4, "Q4")
    sheet.cell(5, 4, "粉质黏土")
    sheet.cell(6, 2, "ZK1")
    sheet.cell(6, 3, 100.0)
    sheet.cell(6, 4, value)
    workbook.save(source)

    with pytest.raises(ValueError, match="有限数字"):
        import_from_table(source, tmp_path / "output")
